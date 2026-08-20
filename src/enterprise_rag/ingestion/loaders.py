"""Robust loaders for common enterprise document formats."""

from __future__ import annotations

import logging
from collections.abc import Callable
from pathlib import Path

from .metadata import default_metadata, parse_front_matter
from .models import DocumentPage, EnterpriseDocument

logger = logging.getLogger(__name__)


def _document_from_text(path: Path, raw_text: str) -> EnterpriseDocument:
    metadata, text = parse_front_matter(raw_text)
    metadata = default_metadata(path, metadata)
    title = str(metadata.get("title", path.stem))
    document_id = str(metadata.get("document_id", path.stem))
    return EnterpriseDocument(
        document_id=document_id,
        text=text,
        title=title,
        source=str(path),
        metadata=metadata,
        pages=[DocumentPage(1, text)],
    )


def load_text(path: Path) -> EnterpriseDocument:
    return _document_from_text(path, path.read_text(encoding="utf-8"))


def load_html(path: Path) -> EnterpriseDocument:
    from bs4 import BeautifulSoup

    raw = path.read_text(encoding="utf-8")
    soup = BeautifulSoup(raw, "html.parser")
    return _document_from_text(path, soup.get_text("\n"))


def load_docx(path: Path) -> EnterpriseDocument:
    from docx import Document

    doc = Document(str(path))
    return _document_from_text(path, "\n\n".join(p.text for p in doc.paragraphs if p.text.strip()))


def load_pdf(path: Path) -> EnterpriseDocument:
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("PDF loading requires the optional 'documents' dependency") from exc
    pdf = fitz.open(str(path))
    pages = [DocumentPage(index + 1, page.get_text("text")) for index, page in enumerate(pdf)]
    metadata, _ = parse_front_matter(pages[0].text if pages else "")
    metadata = default_metadata(path, metadata)
    return EnterpriseDocument(
        document_id=str(metadata["document_id"]),
        text="\n\n".join(page.text for page in pages),
        title=str(metadata["title"]),
        source=str(path),
        metadata=metadata,
        pages=pages,
    )


LOADERS: dict[str, Callable[[Path], EnterpriseDocument]] = {
    ".txt": load_text,
    ".md": load_text,
    ".markdown": load_text,
    ".html": load_html,
    ".htm": load_html,
    ".docx": load_docx,
    ".pdf": load_pdf,
}


def load_file(path: str | Path) -> EnterpriseDocument:
    file_path = Path(path)
    try:
        loader = LOADERS[file_path.suffix.lower()]
    except KeyError as exc:
        raise ValueError(f"Unsupported document format: {file_path.suffix}") from exc
    return loader(file_path)


def load_directory(directory: str | Path) -> tuple[list[EnterpriseDocument], list[dict[str, str]]]:
    """Load supported files while isolating failures to individual documents."""
    root = Path(directory)
    documents: list[EnterpriseDocument] = []
    failures: list[dict[str, str]] = []
    for path in sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in LOADERS):
        try:
            documents.append(load_file(path))
        except Exception as exc:  # document-level isolation is intentional
            logger.exception("Could not parse document %s", path)
            failures.append({"path": str(path), "error": str(exc)})
    return documents, failures

